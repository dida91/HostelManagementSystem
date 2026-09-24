"""Excel (.xlsx) rendering, one sheet per table."""

from __future__ import annotations

import re
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.reports.model import ColumnKind, Report, Table

HEADER_ROW = 4
_HEADER_FILL = PatternFill("solid", fgColor="1F5130")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="D0D5DD")
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_FORMATS: dict[ColumnKind, str] = {
    "money": "#,##0.00",
    "int": "0",
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm",
    "percent": "0.0%",
}


def _sheet_title(title: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub(" ", title).strip()[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(candidate.lower())
    return candidate


def _write(cell: Cell, value: Any, kind: ColumnKind) -> None:
    if value is None or value == "":
        return
    if kind == "text":
        # Always a string cell: user-written text (a complaint, a name) that
        # begins with "=" must never become a live formula in the warden's copy.
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
        cell.data_type = "s"
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        return
    if kind == "datetime" and isinstance(value, datetime):
        # Excel has no timezones; values arrive already in hostel time.
        value = value.replace(tzinfo=None)
    elif kind == "percent" and isinstance(value, int | float | Decimal):
        value = float(value) / 100
    elif kind == "date" and isinstance(value, datetime):
        value = value.date()
    cell.value = value
    cell.number_format = _FORMATS[kind]
    cell.alignment = Alignment(vertical="top")


def _write_table(ws: Any, report: Report, table: Table) -> None:
    ws["A1"] = report.title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{report.subtitle}  |  Generated {report.generated_at:%d %b %Y %H:%M %Z}"
    ws["A2"].font = Font(italic=True, color="555555")

    for i, column in enumerate(table.columns, start=1):
        cell = ws.cell(row=HEADER_ROW, column=i, value=column.header)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=_THIN)
        ws.column_dimensions[get_column_letter(i)].width = max(column.width, 6)

    row_idx = HEADER_ROW
    for row in table.rows:
        row_idx += 1
        for i, (column, value) in enumerate(zip(table.columns, row, strict=True), start=1):
            _write(ws.cell(row=row_idx, column=i), value, column.kind)

    if table.rows:
        last_col = get_column_letter(len(table.columns))
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{row_idx}"
    else:
        row_idx += 1
        ws.cell(row=row_idx, column=1, value="No records.").font = Font(italic=True)
    if table.totals is not None:
        row_idx += 1
        for i, (column, value) in enumerate(zip(table.columns, table.totals, strict=True), 1):
            cell = ws.cell(row=row_idx, column=i)
            _write(cell, value, column.kind)
            cell.font = Font(bold=True)
            cell.border = Border(top=_THIN)

    for note in report.notes:
        row_idx += 2
        ws.cell(row=row_idx, column=1, value=note).font = Font(italic=True, color="555555")
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)


def render_xlsx(report: Report) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for table in report.tables:
        _write_table(workbook.create_sheet(_sheet_title(table.title, used)), report, table)
    workbook.properties.title = report.title
    workbook.properties.creator = "Kutumba Hostel"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
