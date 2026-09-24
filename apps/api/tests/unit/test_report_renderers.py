"""Excel and PDF rendering of the neutral Report model."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.core.config import get_settings
from app.reports import pdf as pdf_module
from app.reports.model import Column, Report, Table
from app.reports.pdf import render_pdf
from app.reports.xlsx import HEADER_ROW, render_xlsx


def _report(rows: list[list[object]], totals: list[object] | None = None) -> Report:
    return Report(
        title="Complaints report",
        subtitle="1 Sep 2026 to 23 Sep 2026",
        generated_at=datetime(2026, 9, 23, 18, 0, tzinfo=ZoneInfo("Asia/Kathmandu")),
        tables=[
            Table(
                title="Complaints: detail/2026",  # characters Excel forbids in sheet names
                columns=[
                    Column("Filed", "date", 11),
                    Column("Summary", width=40),
                    Column("Amount", "money", 12),
                ],
                rows=rows,
                totals=totals,
            )
        ],
        notes=["Resident identities are intentionally omitted from this report."],
    )


def test_xlsx_keeps_user_text_inert_and_numbers_numeric() -> None:
    report = _report(
        [
            [date(2026, 9, 20), '=HYPERLINK("http://evil.example","click")', Decimal("8000")],
            [date(2026, 9, 21), "पानी आएन\x07 (bell char stripped)", Decimal("1234.5")],
        ],
        totals=["Total", None, Decimal("9234.5")],
    )
    ws = load_workbook(BytesIO(render_xlsx(report))).active
    assert ws is not None
    first = HEADER_ROW + 1

    formula_like = ws.cell(row=first, column=2)
    assert formula_like.data_type == "s", "user text must never become a live formula"
    assert formula_like.value.startswith("=HYPERLINK")

    assert ws.cell(row=first + 1, column=2).value == "पानी आएन (bell char stripped)"
    amount = ws.cell(row=first, column=3)
    assert amount.value == 8000 and amount.number_format == "#,##0.00"
    assert ws.cell(row=first + 2, column=1).value == "Total"
    assert ws.cell(row=first + 2, column=3).value == Decimal("9234.5")


def test_xlsx_sheet_titles_are_sanitised() -> None:
    wb = load_workbook(BytesIO(render_xlsx(_report([]))))
    assert wb.sheetnames == ["Complaints  detail 2026"]


def test_xlsx_empty_table_says_so_without_clobbering_totals() -> None:
    ws = load_workbook(BytesIO(render_xlsx(_report([], totals=["Total", None, 0])))).active
    assert ws is not None
    assert ws.cell(row=HEADER_ROW + 1, column=1).value == "No records."
    assert ws.cell(row=HEADER_ROW + 2, column=1).value == "Total"


def test_pdf_renders_and_carries_its_text() -> None:
    content = render_pdf(
        _report([[date(2026, 9, 20), "Room 101 has no water <since> morning & night", 8000]])
    )
    assert content.startswith(b"%PDF")
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
    assert "Complaints report" in text
    assert "<since>" in text and "&" in text, "markup characters are escaped, not parsed"
    assert "8,000.00" in text


def test_pdf_says_when_nepali_cannot_be_drawn(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(get_settings(), "report_devanagari_font", "/nonexistent/font.ttf")
    pdf_module._devanagari_available.cache_clear()
    try:
        content = render_pdf(_report([[date(2026, 9, 20), "पानी आएन", 1]]))
    finally:
        pdf_module._devanagari_available.cache_clear()
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
    assert "Excel export contains the exact text" in text
