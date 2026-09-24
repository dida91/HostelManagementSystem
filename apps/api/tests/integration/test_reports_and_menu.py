"""Report exports over the API, and mess menu management."""

from __future__ import annotations

from io import BytesIO

import pytest
from api_helpers import sita, warden
from httpx import AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.core.db import AsyncSessionLocal
from app.models.user import AuditLog

pytestmark = pytest.mark.asyncio

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.mark.parametrize(
    "path",
    [
        "/api/v1/reports/students",
        "/api/v1/reports/occupancy",
        "/api/v1/reports/fees",
        "/api/v1/reports/complaints",
        "/api/v1/reports/leave",
    ],
)
async def test_every_report_exports_in_both_formats(client: AsyncClient, path: str) -> None:
    staff = await warden(client)

    xlsx = await client.get(f"{path}?format=xlsx", headers=staff)
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.headers["content-type"] == XLSX
    assert xlsx.headers["cache-control"] == "no-store"
    assert "attachment;" in xlsx.headers["content-disposition"]
    workbook = load_workbook(BytesIO(xlsx.content))
    assert workbook.sheetnames

    pdf = await client.get(f"{path}?format=pdf", headers=staff)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert len(PdfReader(BytesIO(pdf.content)).pages) >= 1


async def test_bulk_personal_data_is_warden_only(client: AsyncClient) -> None:
    student = await sita(client)
    for path in ("students", "fees", "occupancy", "complaints"):
        r = await client.get(f"/api/v1/reports/{path}", headers=student)
        assert r.status_code == 403, path


async def test_exports_are_audited(client: AsyncClient) -> None:
    await client.get("/api/v1/reports/occupancy?format=xlsx", headers=await warden(client))
    async with AsyncSessionLocal() as s:
        entries = (
            (
                await s.execute(
                    select(AuditLog).where(
                        AuditLog.action == "report.export", AuditLog.entity_id == "occupancy"
                    )
                )
            )
            .scalars()
            .all()
        )
    assert entries and entries[-1].after["format"] == "xlsx"


async def test_complaints_report_omits_who_complained(client: AsyncClient) -> None:
    student = await sita(client)
    await client.post(
        "/api/v1/complaints",
        headers=student,
        json={"text": "The corridor light near room 102 flickers all night."},
    )
    r = await client.get("/api/v1/reports/complaints?format=xlsx", headers=await warden(client))
    cells = [
        str(c.value)
        for ws in load_workbook(BytesIO(r.content)).worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    ]
    text = " ".join(cells)
    assert "corridor light" in text
    assert "KH-2026-001" not in text and "Sita" not in text


async def test_report_date_range_is_validated(client: AsyncClient) -> None:
    staff = await warden(client)
    inverted = await client.get(
        "/api/v1/reports/leave?date_from=2026-09-10&date_to=2026-09-01", headers=staff
    )
    assert inverted.status_code == 422
    too_long = await client.get(
        "/api/v1/reports/complaints?date_from=2024-01-01&date_to=2026-01-01", headers=staff
    )
    assert too_long.status_code == 422


async def test_menu_can_be_set_and_removed_by_staff(client: AsyncClient) -> None:
    staff = await warden(client)
    r = await client.put(
        "/api/v1/mess/menu/6/SNACKS",
        headers=staff,
        json={"items": "Chiura, aloo sadeko", "serving_time": "16:00-17:00"},
    )
    assert r.status_code == 200, r.text
    assert r.json() == {
        "day_of_week": 6,
        "meal_type": "SNACKS",
        "items": "Chiura, aloo sadeko",
        "serving_time": "16:00-17:00",
    }

    replaced = await client.put(
        "/api/v1/mess/menu/6/SNACKS", headers=staff, json={"items": "Sel roti"}
    )
    assert replaced.json()["items"] == "Sel roti" and replaced.json()["serving_time"] is None

    saturday = (await client.get("/api/v1/mess/menu?day_of_week=6", headers=staff)).json()
    assert [m["items"] for m in saturday if m["meal_type"] == "SNACKS"] == ["Sel roti"]

    assert (await client.delete("/api/v1/mess/menu/6/SNACKS", headers=staff)).status_code == 204
    assert (await client.delete("/api/v1/mess/menu/6/SNACKS", headers=staff)).status_code == 404


async def test_students_cannot_change_the_menu(client: AsyncClient) -> None:
    r = await client.put(
        "/api/v1/mess/menu/1/LUNCH", headers=await sita(client), json={"items": "Pizza"}
    )
    assert r.status_code == 403


async def test_menu_day_must_exist(client: AsyncClient) -> None:
    r = await client.put(
        "/api/v1/mess/menu/7/LUNCH", headers=await warden(client), json={"items": "Dal bhat"}
    )
    assert r.status_code == 422
